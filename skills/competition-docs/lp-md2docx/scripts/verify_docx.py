#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""重新打开 docx，校验样式表、多级列表、标题层级与图片。"""

import argparse
import sys

from docx import Document
from docx.oxml.ns import qn


HEADING_STYLE_NAMES = {
    1: "一级标题",
    2: "二级标题",
    3: "三级标题",
    4: "四级标题",
    5: "五级标题",
    6: "六级标题",
}


def _find_numbering_part(document):
    for part in document.part.package.iter_parts():
        if part.partname.endswith("/word/numbering.xml"):
            return part
    return None


def main(argv=None):
    ap = argparse.ArgumentParser(description="校验生成的 docx")
    ap.add_argument("docx", help="要校验的 .docx 文件")
    args = ap.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    document = Document(args.docx)
    problems = []

    # 1) 标题样式：中文字体、outline 层级、编号
    for level in range(1, 7):
        name = HEADING_STYLE_NAMES[level]
        try:
            style = document.styles[name]
        except KeyError:
            problems.append(f"缺少样式 {name}")
            continue
        ppr = style.element.pPr
        ol = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        rpr = style.element.rPr
        rf = rpr.rFonts if rpr is not None else None
        east = rf.get(qn("w:eastAsia")) if rf is not None else None
        size = style.font.size.pt if style.font.size else None
        ol_val = ol.get(qn("w:val")) if ol is not None else None
        print(
            f"样式 {name}: 中文字体={east} 字号={size}pt "
            f"outline={ol_val} 编号={'有' if numpr is not None else '无'}"
        )
        if level <= 4 and (ol is None or numpr is None):
            problems.append(f"{name} 缺少 outlineLvl 或 numPr")

    title_style = document.styles["Title"]
    tppr = title_style.element.pPr
    tjc = tppr.find(qn("w:jc")) if tppr is not None else None
    print(
        f"样式 Title: 对齐="
        f"{tjc.get(qn('w:val')) if tjc is not None else '继承'} "
        f"编号={'有' if (tppr is not None and tppr.find(qn('w:numPr')) is not None) else '无'}"
    )
    if tppr is not None and tppr.find(qn("w:numPr")) is not None:
        problems.append("Title 大标题不应带编号")

    # 2) 多级列表定义（只检查我们写入的那组）
    numbering = _find_numbering_part(document)
    if numbering is None:
        problems.append("缺少 numbering.xml（多级列表未生效）")
    else:
        heading_style = document.styles[HEADING_STYLE_NAMES[1]]
        numpr = heading_style.element.pPr.find(qn("w:numPr"))
        num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
        abstract_id = None
        for n in numbering.element.findall(qn("w:num")):
            if n.get(qn("w:numId")) == num_id:
                ref = n.find(qn("w:abstractNumId"))
                abstract_id = ref.get(qn("w:val")) if ref is not None else None
        print(f"\n多级列表定义（numId={num_id} abstractId={abstract_id}）:")
        for absn in numbering.element.findall(qn("w:abstractNum")):
            if absn.get(qn("w:abstractNumId")) != abstract_id:
                continue
            for lvl in absn.findall(qn("w:lvl")):
                txt = lvl.find(qn("w:lvlText"))
                fmt = lvl.find(qn("w:numFmt"))
                print(
                    f"  ilvl={lvl.get(qn('w:ilvl'))} "
                    f"格式={fmt.get(qn('w:val')) if fmt is not None else None} "
                    f"文本={txt.get(qn('w:val')) if txt is not None else None}"
                )

    # 3) 文档标题序列
    print("\n文档标题序列:")
    for p in document.paragraphs:
        if p.style.name in HEADING_STYLE_NAMES.values() or p.style.name == "Title":
            print(f"  [{p.style.name}] {p.text}")

    # 4) 注入的自定义样式（Word 样式面板可见）
    print("\n注入样式检查:")
    style_names = [s.name for s in document.styles]
    for name, stype in (
        ("Title", "paragraph"),
        ("论文正文", "paragraph"),
        ("图注", "paragraph"),
        ("表注", "paragraph"),
        ("表内文字", "paragraph"),
    ):
        style = document.styles[name] if name in style_names else None
        if style is None:
            problems.append(f"缺少注入样式 {name}")
            print(f"  {name}: 缺失")
            continue
        el = style.element
        up = el.find(qn("w:uiPriority"))
        qf = el.find(qn("w:qFormat"))
        print(
            f"  {name}（{stype}）: uiPriority={up.get(qn('w:val')) if up is not None else '无'} "
            f"qFormat={'有' if qf is not None else '无'}"
        )
        if up is None or qf is None:
            problems.append(f"样式 {name} 未标记为可见（缺 uiPriority/qFormat）")

    table_styles_found = [
        s.name
        for s in document.styles
        if s.element.get(qn("w:type")) == "table"
        and s.name in ("三线表", "标准表格")
    ]
    print(f"  表格样式: {table_styles_found if table_styles_found else '缺失'}")
    if not table_styles_found:
        problems.append("缺少注入的表格样式（三线表/标准表格）")

    body_paras = [p for p in document.paragraphs if p.style.name == "论文正文"]
    print(f"  论文正文段落数: {len(body_paras)}")
    if not body_paras:
        problems.append("正文段落未使用论文正文样式")

    # 图注/表注自动编号（图1/表1）
    for name in ("图注", "表注"):
        style = document.styles[name]
        ppr = style.element.pPr
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        has = numpr is not None
        print(f"  {name} 自动编号: {'有' if has else '无'}")
        if not has:
            problems.append(f"{name} 缺少自动编号（图1/表1）")

    # 5) 原生列表
    list_paras = [p for p in document.paragraphs if p.style.name == "List Paragraph"]
    print(f"\n原生列表段落: {len(list_paras)}")
    for p in list_paras[:4]:
        numpr = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
        if numpr is None:
            problems.append("存在未挂编号的 List Paragraph 段落")
            continue
        num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
        ilvl = numpr.find(qn("w:ilvl")).get(qn("w:val"))
        print(f"  numId={num_id} ilvl={ilvl} | {p.text}")
    if list_paras and any(
        p._p.pPr is None or p._p.pPr.find(qn("w:numPr")) is None for p in list_paras
    ):
        problems.append("原生列表编号未完整挂载")

    # 6) 表格
    print(f"\n表格数量: {len(document.tables)}")
    known_table_styles = {"三线表", "标准表格"}
    for t in document.tables:
        print(f"  样式={t.style.name if t.style else '无'} 尺寸={len(t.rows)}x{len(t.columns)}")
        if t.style is None or t.style.name not in known_table_styles:
            problems.append(f"表格使用了未知样式: {t.style.name if t.style else '无'}")

    # 7) 图片
    print(f"\n内嵌图片数量: {len(document.inline_shapes)}")

    if problems:
        print("\n[失败] " + "; ".join(problems))
        return 1
    print("\n[通过] 样式表 / 多级列表 / 标题层级检查通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
